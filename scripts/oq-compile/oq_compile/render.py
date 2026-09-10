"""Write the deliverables: a deterministic CSV extract and the workbook."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from .manifest import Manifest
from .pivot import FAIL, NOT_RUN, PASS


@dataclass(frozen=True)
class Provenance:
    primary_commit: str
    associate_commits: tuple[str, ...]
    elspais_version: str
    tool_version: str
    scope_name: str
    scope_lines: tuple[str, ...]
    generated_at: str
    req_count: int
    uat_count: int


def _headers(columns: tuple[str, ...], width: int) -> list[str]:
    """Extend the declared header to the widest row.

    The last declared column repeats: rows carry a variable number of
    journeys (or requirements), each in its own column, and every one of them
    is the same kind of thing.
    """
    if not columns:
        return [""] * width
    header = list(columns)
    while len(header) < width:
        header.append(columns[-1])
    return header[:width] if width else header


def write_csv(path: Path, header: tuple[str, ...], rows: list[list[str]]) -> None:
    """Write the extract that is committed and diffed.

    Every row is padded to the widest, so the file is a rectangle and a diff
    reports a changed cell rather than a reshaped row. Newlines are fixed to
    "\\n" so the bytes do not depend on the platform that produced them.
    """
    width = max([len(header), *(len(r) for r in rows)] or [len(header)])
    with Path(path).open("w", newline="\n", encoding="utf-8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(_headers(header, width))
        for row in rows:
            writer.writerow(list(row) + [""] * (width - len(row)))


def _provenance_rows(manifest: Manifest, prov: Provenance) -> list[list[str]]:
    return [
        ["Report", manifest.title],
        ["Project", manifest.project],
        ["Generated at (UTC)", prov.generated_at],
        [],
        ["Scope name", prov.scope_name],
        *[["Scope", line] for line in prov.scope_lines],
        ["Requirements reported", str(prov.req_count)],
        ["UAT test cases reported", str(prov.uat_count)],
        [],
        ["Primary repository commit", prov.primary_commit],
        *[["Associate repository commit", c] for c in prov.associate_commits],
        ["elspais version", prov.elspais_version],
        ["Generator version", prov.tool_version],
        [],
        ["Verdict legend", ""],
        [PASS, "Every assertion the requirement expects is verified, and no validating journey failed."],
        [FAIL, "At least one validating journey failed."],
        [NOT_RUN, "No validating journey has been run, or coverage is partial. Not a failure."],
        [],
        [
            "Note",
            "Coverage figures here are federated across every repository in the "
            "graph. The checks report aggregates only its own repository, so the "
            "two surfaces answer different questions and need not agree.",
        ],
    ]


def write_workbook(
    path: Path,
    manifest: Manifest,
    req_rows_: list[list[str]],
    uat_rows_: list[list[str]],
    provenance: Provenance,
) -> None:
    """Write the three-sheet workbook.

    Not committed: a workbook embeds creation and modification times and
    per-entry archive timestamps, so two with identical content differ byte
    for byte. It is uploaded as an artifact bound to the run instead.
    """
    wb = Workbook()

    req = wb.active
    req.title = manifest.req_sheet.name
    req_width = max([len(manifest.req_sheet.columns), *(len(r) for r in req_rows_)] or [0])
    req.append(_headers(manifest.req_sheet.columns, req_width))
    for row in req_rows_:
        req.append(row)

    uat = wb.create_sheet(manifest.uat_sheet.name)
    uat_width = max([len(manifest.uat_sheet.columns), *(len(r) for r in uat_rows_)] or [0])
    uat.append(_headers(manifest.uat_sheet.columns, uat_width))
    for row in uat_rows_:
        uat.append(row)

    prov = wb.create_sheet(manifest.provenance_sheet_name)
    for row in _provenance_rows(manifest, provenance):
        prov.append(row)

    wb.save(Path(path))

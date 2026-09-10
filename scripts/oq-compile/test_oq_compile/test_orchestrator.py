from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

SCRIPT = Path(__file__).resolve().parents[1] / "compile-oq.py"


@pytest.fixture(scope="module")
def compile_oq():
    sys.path.insert(0, str(SCRIPT.parent))
    spec = importlib.util.spec_from_file_location("compile_oq", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_build_writes_all_three_outputs(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_path, sample_graph_path
):
    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    assert counts == (3, 3)
    assert (tmp_path / "reports" / "oq-req.csv").exists()
    assert (tmp_path / "reports" / "oq-uat.csv").exists()
    assert (tmp_path / "build" / "oq.xlsx").exists()


def test_csv_outputs_are_reproducible(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_path, sample_graph_path
):
    import hashlib

    digests = []
    for run in ("one", "two"):
        target = tmp_path / run
        compile_oq.build(
            manifest_path=sample_manifest_path,
            trace_path=sample_trace_path,
            graph_path=sample_graph_path,
            out_csv_dir=target,
            out_xlsx=target / "oq.xlsx",
            provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        )
        digests.append(
            (
                hashlib.sha256((target / "oq-req.csv").read_bytes()).hexdigest(),
                hashlib.sha256((target / "oq-uat.csv").read_bytes()).hexdigest(),
            )
        )
    assert digests[0] == digests[1]


def test_generated_at_override_reaches_the_workbook(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_path, sample_graph_path
):
    """provenance_overrides is asserted as an input by every other test here;
    this asserts it as an effect. If build() silently ignored the override
    and stamped the real clock instead, every other test would still pass."""
    import openpyxl

    out_xlsx = tmp_path / "build" / "oq.xlsx"
    compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=out_xlsx,
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    manifest = compile_oq.Manifest.from_path(sample_manifest_path)
    wb = openpyxl.load_workbook(out_xlsx)
    prov_rows = dict(
        (row[0], row[1])
        for row in wb[manifest.provenance_sheet_name].iter_rows(values_only=True)
        if row and row[0]
    )
    assert prov_rows["Generated at (UTC)"] == "2026-09-09T12:00:00Z"


def test_empty_selection_refused_by_default(
    compile_oq, tmp_path, sample_manifest_path, sample_graph_path
):
    # A trace file selecting zero requirements is exactly what the CUR-1925
    # loader bug produced for every scoped run: a well-formed, correctly-
    # provenanced, entirely empty report at exit 0. build() must refuse it.
    empty_trace = tmp_path / "trace-empty.json"
    empty_trace.write_text("[]")

    with pytest.raises(ValueError) as excinfo:
        compile_oq.build(
            manifest_path=sample_manifest_path,
            trace_path=empty_trace,
            graph_path=sample_graph_path,
            out_csv_dir=tmp_path / "reports",
            out_xlsx=tmp_path / "build" / "oq.xlsx",
            provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        )
    assert "example-scope" in str(excinfo.value)
    assert "--allow-empty" in str(excinfo.value)
    # Refused before any deliverable is written.
    assert not (tmp_path / "reports" / "oq-req.csv").exists()
    assert not (tmp_path / "build" / "oq.xlsx").exists()


def test_empty_selection_permitted_with_flag(
    compile_oq, tmp_path, sample_manifest_path, sample_graph_path
):
    empty_trace = tmp_path / "trace-empty.json"
    empty_trace.write_text("[]")

    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=empty_trace,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        allow_empty=True,
    )
    assert counts == (0, 0)
    assert (tmp_path / "reports" / "oq-req.csv").exists()
    assert (tmp_path / "build" / "oq.xlsx").exists()


def test_non_empty_selection_unaffected_by_the_empty_check(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_path, sample_graph_path
):
    # allow_empty defaults to False; a non-empty selection must not trip
    # the new check.
    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    assert counts == (3, 3)


def test_main_prints_the_stdout_contract(
    compile_oq,
    tmp_path,
    monkeypatch,
    capsys,
    sample_manifest_path,
    sample_trace_path,
    sample_graph_path,
):
    """Task 7's composite action parses this exact printed line with `sed`
    to produce its two GitHub Actions outputs; a change to the format here
    would pass every other test in this file and break that step silently."""
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "compile-oq.py",
            "--manifest",
            str(sample_manifest_path),
            "--trace",
            str(sample_trace_path),
            "--graph",
            str(sample_graph_path),
            "--out-csv-dir",
            str(tmp_path / "reports"),
            "--out-xlsx",
            str(tmp_path / "build" / "oq.xlsx"),
        ],
    )
    exit_code = compile_oq.main()
    captured = capsys.readouterr()
    assert exit_code == 0
    assert captured.out == "OQ report: 3 requirements, 3 UAT test cases\n"


def test_scoped_trace_fixture_drives_the_orchestrator(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_scoped_path, sample_graph_path
):
    """Every other test here feeds the bare-list trace shape, which no scoped
    production run emits. A scoped manifest -- the normal case -- makes the
    exporter emit a dict with the rows under `nodes`, and reading that key
    wrongly is the exact defect that rendered an empty report at exit 0. Drive
    the orchestrator end to end on the dict shape so the wiring, not just the
    loader, is exercised against it.

    `JNY-FIX-01` is deliberately absent from the graph fixture, so its
    Description renders blank: one missing title is a gap in the graph, not the
    total loader failure the zero-journey guard exists for.
    """
    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=sample_trace_scoped_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    assert counts == (1, 1)
    req_csv = (tmp_path / "reports" / "oq-req.csv").read_text()
    assert "SPN-PRD-fixture-obligation" in req_csv
    assert "UAT-JNY-FIX-01" in req_csv
    uat_csv = (tmp_path / "reports" / "oq-uat.csv").read_text()
    assert "UAT-JNY-FIX-01" in uat_csv


def _manifest_with(tmp_path, sample_manifest_dict, **extra):
    import yaml

    raw = dict(sample_manifest_dict)
    raw.update(extra)
    path = tmp_path / "manifest.yaml"
    path.write_text(yaml.safe_dump(raw))
    return path


def test_declared_namespace_missing_is_refused(
    compile_oq,
    tmp_path,
    sample_manifest_dict,
    sample_trace_path,
    sample_graph_path,
):
    """The 80%-incomplete report: a federated associate is simply not
    configured, so its namespace contributes no row and everything else about
    the run looks healthy."""
    manifest = _manifest_with(
        tmp_path, sample_manifest_dict, require_namespaces=["SPN", "PLT"]
    )
    with pytest.raises(ValueError) as excinfo:
        compile_oq.build(
            manifest_path=manifest,
            trace_path=sample_trace_path,
            graph_path=sample_graph_path,
            out_csv_dir=tmp_path / "reports",
            out_xlsx=tmp_path / "build" / "oq.xlsx",
            provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        )
    message = str(excinfo.value)
    assert "PLT" in message      # what is missing
    assert "SPN" in message      # what was found
    # Refused before any deliverable is written.
    assert not (tmp_path / "reports" / "oq-req.csv").exists()
    assert not (tmp_path / "build" / "oq.xlsx").exists()


def test_declared_namespace_missing_is_refused_even_with_allow_empty(
    compile_oq,
    tmp_path,
    sample_manifest_dict,
    sample_trace_path,
    sample_graph_path,
):
    """--allow-empty says an empty report may be honest. It does not say a
    report missing a namespace the manifest declares may be."""
    manifest = _manifest_with(
        tmp_path, sample_manifest_dict, require_namespaces=["PLT"]
    )
    with pytest.raises(ValueError):
        compile_oq.build(
            manifest_path=manifest,
            trace_path=sample_trace_path,
            graph_path=sample_graph_path,
            out_csv_dir=tmp_path / "reports",
            out_xlsx=tmp_path / "build" / "oq.xlsx",
            allow_empty=True,
        )


def test_all_declared_namespaces_present_passes(
    compile_oq,
    tmp_path,
    sample_manifest_dict,
    sample_trace_path,
    sample_graph_path,
):
    manifest = _manifest_with(
        tmp_path, sample_manifest_dict, require_namespaces=["SPN"]
    )
    counts = compile_oq.build(
        manifest_path=manifest,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    assert counts == (3, 3)


def _urs_backed_manifest_path(tmp_path, primary_root, *, columns_extra=True):
    """An OQ manifest declaring a urs_manifest, alongside a minimal (but
    valid) URS manifest under primary_root -- everything resolve_sponsor_name
    needs to look beside it for sponsor-info.yaml."""
    import yaml

    urs_dir = primary_root / "spec" / "URS-manifest"
    urs_dir.mkdir(parents=True, exist_ok=True)
    (urs_dir / "urs.yaml").write_text("document: {}\n")

    req_columns = ["Req ID", "Description", "Test Result", "UAT Result", "UAT Test Case ID"]
    if columns_extra:
        req_columns.insert(1, "URS Section")
    doc = {
        "document": {"title": "OQ", "project": "Example Study"},
        "scope": "example-scope",
        "urs_manifest": "spec/URS-manifest/urs.yaml",
        "sheets": {
            "req": {"name": "Requirements", "columns": req_columns},
            "uat": {
                "name": "UAT Test Cases",
                "columns": [
                    "UAT Test Case ID", "Description", "Pass/Fail",
                    "User Journey ID", "Req ID",
                ],
            },
        },
    }
    path = tmp_path / "oq.yaml"
    path.write_text(yaml.safe_dump(doc))
    return path, urs_dir


def test_sponsor_row_resolved_from_sponsor_info_yaml(
    compile_oq, tmp_path, sample_trace_path, sample_graph_path
):
    """End-to-end: a real sponsor-info.yaml beside the URS manifest this OQ
    manifest declares is the only source for the Sponsor row's value."""
    import openpyxl

    from oq_compile.render import VENDOR_NAME

    primary_root = tmp_path / "primary"
    manifest_path, urs_dir = _urs_backed_manifest_path(tmp_path, primary_root)
    (urs_dir / "sponsor-info.yaml").write_text(
        "sponsor_name: Example Sponsor, Inc.\n"
    )
    out_xlsx = tmp_path / "build" / "oq.xlsx"

    compile_oq.build(
        manifest_path=manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=out_xlsx,
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        primary_root=primary_root,
    )

    wb = openpyxl.load_workbook(out_xlsx)
    prov_rows = dict(
        (row[0], row[1])
        for row in wb["Provenance"].iter_rows(values_only=True)
        if row and row[0]
    )
    assert prov_rows["Vendor"] == VENDOR_NAME
    assert prov_rows["Sponsor"] == "Example Sponsor, Inc."


def test_sponsor_row_absent_without_a_sponsor_info_file(
    compile_oq, tmp_path, sample_trace_path, sample_graph_path
):
    """A urs_manifest is declared, but no sponsor-info.yaml sits beside it
    -- the report is still produced, just without a Sponsor row."""
    import openpyxl

    from oq_compile.render import VENDOR_NAME

    primary_root = tmp_path / "primary"
    manifest_path, _urs_dir = _urs_backed_manifest_path(tmp_path, primary_root)
    out_xlsx = tmp_path / "build" / "oq.xlsx"

    compile_oq.build(
        manifest_path=manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=out_xlsx,
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        primary_root=primary_root,
    )

    wb = openpyxl.load_workbook(out_xlsx)
    prov_rows = dict(
        (row[0], row[1])
        for row in wb["Provenance"].iter_rows(values_only=True)
        if row and row[0]
    )
    assert prov_rows["Vendor"] == VENDOR_NAME
    assert "Sponsor" not in prov_rows


def test_no_namespace_declaration_behaves_as_before(
    compile_oq, tmp_path, sample_manifest_path, sample_trace_path, sample_graph_path
):
    """The stock fixture manifest declares no require_namespaces, so the
    generator stays agnostic about who is federated."""
    manifest = compile_oq.Manifest.from_path(sample_manifest_path)
    assert manifest.require_namespaces == ()
    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
    )
    assert counts == (3, 3)


def test_zero_uat_rows_refused_by_default(
    compile_oq, tmp_path, sample_manifest_path, sample_graph_path
):
    """A full REQ sheet with an empty UAT sheet: what a dropped or renamed
    `journeys` key upstream produces. Counting requirements alone passes it."""
    import json

    trace = tmp_path / "trace.json"
    trace.write_text(
        json.dumps(
            [
                {
                    "id": "SPN-PRD-audit-log",
                    "title": "Audit Log",
                    "level": "PRD",
                    "status": "Active",
                    "uat_verified": {"ratio": 0.0},
                    "verified": {"ratio": 0.0, "carried": False},
                    "tested": {"failed": 0.0},
                    "journeys": [],
                }
            ]
        )
    )
    with pytest.raises(ValueError) as excinfo:
        compile_oq.build(
            manifest_path=sample_manifest_path,
            trace_path=trace,
            graph_path=sample_graph_path,
            out_csv_dir=tmp_path / "reports",
            out_xlsx=tmp_path / "build" / "oq.xlsx",
        )
    assert "UAT" in str(excinfo.value)
    assert "--allow-empty" in str(excinfo.value)
    assert not (tmp_path / "reports" / "oq-req.csv").exists()
    assert not (tmp_path / "build" / "oq.xlsx").exists()


def test_zero_uat_rows_permitted_with_flag(
    compile_oq, tmp_path, sample_manifest_path, sample_graph_path
):
    import json

    trace = tmp_path / "trace.json"
    trace.write_text(
        json.dumps(
            [
                {
                    "id": "SPN-PRD-audit-log",
                    "title": "Audit Log",
                    "level": "PRD",
                    "status": "Active",
                    "uat_verified": {"ratio": 0.0},
                    "verified": {"ratio": 0.0, "carried": False},
                    "tested": {"failed": 0.0},
                    "journeys": [],
                }
            ]
        )
    )
    counts = compile_oq.build(
        manifest_path=sample_manifest_path,
        trace_path=trace,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=tmp_path / "build" / "oq.xlsx",
        allow_empty=True,
    )
    assert counts == (1, 0)


def test_protocol_rows_resolved_from_sponsor_info_yaml(
    compile_oq, tmp_path, sample_trace_path, sample_graph_path
):
    """End-to-end: a real sponsor-info.yaml beside the URS manifest this OQ
    manifest declares is the only source for the Protocol number and
    Protocol version rows, same as the Sponsor row above."""
    import openpyxl

    primary_root = tmp_path / "primary"
    manifest_path, urs_dir = _urs_backed_manifest_path(tmp_path, primary_root)
    (urs_dir / "sponsor-info.yaml").write_text(
        "protocol_number: EXAMPLE-0001\nprotocol_version: '3.0'\n"
    )
    out_xlsx = tmp_path / "build" / "oq.xlsx"

    compile_oq.build(
        manifest_path=manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=out_xlsx,
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        primary_root=primary_root,
    )

    wb = openpyxl.load_workbook(out_xlsx)
    prov_rows = dict(
        (row[0], row[1])
        for row in wb["Provenance"].iter_rows(values_only=True)
        if row and row[0]
    )
    assert prov_rows["Protocol number"] == "EXAMPLE-0001"
    assert prov_rows["Protocol version"] == "3.0"


def test_protocol_rows_absent_without_a_sponsor_info_file(
    compile_oq, tmp_path, sample_trace_path, sample_graph_path
):
    """A urs_manifest is declared, but no sponsor-info.yaml sits beside it
    -- the report is still produced, just without Protocol number or
    Protocol version rows (same file-absent case the Sponsor row covers)."""
    import openpyxl

    primary_root = tmp_path / "primary"
    manifest_path, _urs_dir = _urs_backed_manifest_path(tmp_path, primary_root)
    out_xlsx = tmp_path / "build" / "oq.xlsx"

    compile_oq.build(
        manifest_path=manifest_path,
        trace_path=sample_trace_path,
        graph_path=sample_graph_path,
        out_csv_dir=tmp_path / "reports",
        out_xlsx=out_xlsx,
        provenance_overrides={"generated_at": "2026-09-09T12:00:00Z"},
        primary_root=primary_root,
    )

    wb = openpyxl.load_workbook(out_xlsx)
    prov_rows = dict(
        (row[0], row[1])
        for row in wb["Provenance"].iter_rows(values_only=True)
        if row and row[0]
    )
    assert "Protocol number" not in prov_rows
    assert "Protocol version" not in prov_rows

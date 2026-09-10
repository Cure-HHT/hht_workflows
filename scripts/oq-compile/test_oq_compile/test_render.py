from __future__ import annotations

import hashlib

import openpyxl

from oq_compile.manifest import Manifest
from oq_compile.render import Provenance, write_csv, write_workbook


def _provenance():
    return Provenance(
        primary_commit="abc1234",
        associate_commits=("def5678",),
        elspais_version="0.123.0",
        tool_version="v1.0.0",
        scope_name="example-scope",
        scope_lines=("Scope: level PRD or GUI",),
        generated_at="2026-09-09T12:00:00Z",
        req_count=3,
        uat_count=3,
    )


def test_csv_pads_rows_to_the_widest(tmp_path):
    out = tmp_path / "req.csv"
    write_csv(out, ("A", "B", "C"), [["1", "2", "3", "4", "5"], ["6", "7", "8"]])
    lines = out.read_text().splitlines()
    assert lines[0] == "A,B,C,C,C"
    assert lines[2] == "6,7,8,,"


def test_csv_is_byte_identical_across_runs(tmp_path):
    rows = [["1", "2", "3", "4"]]
    a, b = tmp_path / "a.csv", tmp_path / "b.csv"
    write_csv(a, ("A", "B", "C"), rows)
    write_csv(b, ("A", "B", "C"), rows)
    assert hashlib.sha256(a.read_bytes()).hexdigest() == hashlib.sha256(
        b.read_bytes()
    ).hexdigest()


def test_workbook_has_three_named_sheets_provenance_first(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [["SPN-PRD-a", "A", "PASS", "PASS", "JNY-1"]],
                   [["UAT-JNY-1", "J", "PASS", "JNY-1", "SPN-PRD-a"]], _provenance())
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Provenance", "REQ", "UAT Test Cases"]
    assert wb.active.title == "Provenance"


def test_workbook_writes_headers_and_rows(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [["SPN-PRD-a", "A", "PASS", "NOT RUN", "JNY-1"]],
                   [["UAT-JNY-1", "J", "PASS", "JNY-1", "SPN-PRD-a"]], _provenance())
    ws = openpyxl.load_workbook(out)["REQ"]
    assert [c.value for c in ws[1]][:5] == [
        "Req ID", "Description", "Test Result", "UAT Result", "UAT Test Case ID",
    ]
    assert [c.value for c in ws[2]] == [
        "SPN-PRD-a", "A", "PASS", "NOT RUN", "JNY-1",
    ]


def test_provenance_sheet_states_scope_and_timestamp(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [], [], _provenance())
    ws = openpyxl.load_workbook(out)["Provenance"]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert "example-scope" in text
    assert "2026-09-09T12:00:00Z" in text
    assert "0.123.0" in text
    assert "NOT RUN" in text


def test_extra_columns_beyond_the_declared_header(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m,
                   [["SPN-PRD-a", "A", "PASS", "PASS", "JNY-1", "JNY-2", "JNY-3"]],
                   [], _provenance())
    ws = openpyxl.load_workbook(out)["REQ"]
    assert [c.value for c in ws[1]][:7] == [
        "Req ID", "Description", "Test Result", "UAT Result",
        "UAT Test Case ID", "UAT Test Case ID", "UAT Test Case ID",
    ]


def _provenance_text(tmp_path, manifest_path):
    m = Manifest.from_path(manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [], [], _provenance())
    ws = openpyxl.load_workbook(out)[m.provenance_sheet_name]
    return "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )


def test_legend_explains_both_verdict_columns(tmp_path, sample_manifest_path):
    """A legend explaining only one column leaves a regulator unable to read
    the other. Both column titles must appear, named as the manifest declares
    them rather than hardcoded."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    assert "Test Result" in text
    assert "UAT Result" in text


def test_legend_uses_the_manifest_column_titles(tmp_path, sample_manifest_dict):
    """A consumer that renames the verdict columns gets a legend matching its
    own sheet; nothing about the titles is baked into the generator."""
    import yaml

    raw = dict(sample_manifest_dict)
    raw["sheets"] = {
        **raw["sheets"],
        "req": {
            "name": "REQ",
            "columns": [
                "Req ID", "Description", "Verification", "Acceptance",
                "UAT Test Case ID",
            ],
        },
    }
    path = tmp_path / "manifest.yaml"
    path.write_text(yaml.safe_dump(raw))
    text = _provenance_text(tmp_path, path)
    assert "Verification" in text
    assert "Acceptance" in text
    assert "Test Result" not in text


def test_column_definitions_use_the_manifest_uat_column_titles(
    tmp_path, sample_manifest_dict
):
    """A consumer that renames a UAT-sheet column gets column definitions
    matching its own sheet, the same guarantee as the REQ-sheet titles."""
    import yaml

    raw = dict(sample_manifest_dict)
    raw["sheets"] = {
        **raw["sheets"],
        "uat": {
            "name": "UAT Test Cases",
            "columns": [
                "UAT Test Case ID", "Description", "Outcome",
                "User Journey ID", "Req ID",
            ],
        },
    }
    path = tmp_path / "manifest.yaml"
    path.write_text(yaml.safe_dump(raw))
    text = _provenance_text(tmp_path, path)
    assert "Outcome" in text
    assert "Pass/Fail" not in text


def test_provenance_sheet_reads_facts_then_definitions_then_legend(
    tmp_path, sample_manifest_path
):
    """The sheet is laid out so it reads well as the first thing someone
    sees: identifying facts first, then the column definitions, then the
    verdict legend."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    facts_at = text.index("Generated at (UTC)")
    definitions_at = text.index("Column definitions")
    legend_at = text.index("Verdict legend")
    assert facts_at < definitions_at < legend_at


def test_legend_distinguishes_the_two_not_run_meanings(
    tmp_path, sample_manifest_path
):
    """NOT RUN means a different absence in each column -- no test result
    ingested, versus no validating journey run. A legend stating one
    definition for both would mislead."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    assert "no test result has been ingested" in text.lower()
    assert "no validating journey has been run" in text.lower()


def test_legend_explains_the_carried_marker(tmp_path, sample_manifest_path):
    """The marker appears in report cells, so its meaning must be stated."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    assert "(carried)" in text
    assert "carried forward from a baseline" in text.lower()


def test_stale_federation_note_is_gone(tmp_path, sample_manifest_path):
    """The report no longer states any coverage figure -- only verdicts --
    so a note contrasting its coverage numbers with the checks report's own
    aggregate describes something not on the sheet."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    assert "federated across every repository" not in text
    assert "checks report aggregates only its own repository" not in text


def test_provenance_sheet_defines_every_grid_column(tmp_path, sample_manifest_path):
    """A reader must be able to learn what every column means from the
    workbook alone, named as the manifest declares them."""
    text = _provenance_text(tmp_path, sample_manifest_path)
    for title in ("Req ID", "Description", "Test Result", "UAT Result", "UAT Test Case ID"):
        assert title in text
    for title in ("Pass/Fail", "User Journey ID"):
        assert title in text
    assert "validating test-case identifier" in text.lower()
    assert "validated requirement" in text.lower() or "requirement id this test case validates" in text.lower()


def test_grid_sheet_headers_are_bold(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [["SPN-PRD-a", "A", "PASS", "PASS", "JNY-1"]],
                   [["UAT-JNY-1", "J", "PASS", "JNY-1", "SPN-PRD-a"]], _provenance())
    wb = openpyxl.load_workbook(out)
    for name in ("REQ", "UAT Test Cases"):
        ws = wb[name]
        assert all(c.font.bold for c in ws[1])
        # A data row is not bold.
        assert not any(c.font.bold for c in ws[2])


def test_provenance_labels_are_bold_but_prose_is_not(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [], [], _provenance())
    ws = openpyxl.load_workbook(out)["Provenance"]
    labeled = [row for row in ws.iter_rows() if row[0].value]
    assert labeled  # sanity: there are labeled rows
    assert all(row[0].font.bold for row in labeled)
    # A row whose first cell is blank carries explanatory prose, not a label.
    prose_rows = [row for row in ws.iter_rows() if not row[0].value and row[1].value]
    assert prose_rows
    assert not any(row[0].font.bold for row in prose_rows)


def test_cells_wrap_on_all_three_sheets(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    write_workbook(out, m, [["SPN-PRD-a", "A", "PASS", "PASS", "JNY-1"]],
                   [["UAT-JNY-1", "J", "PASS", "JNY-1", "SPN-PRD-a"]], _provenance())
    wb = openpyxl.load_workbook(out)
    for name in ("Provenance", "REQ", "UAT Test Cases"):
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                assert cell.alignment.wrap_text is True


def test_column_widths_are_bounded_and_content_driven(tmp_path, sample_manifest_path):
    m = Manifest.from_path(sample_manifest_path)
    out = tmp_path / "oq.xlsx"
    long_title = "A" * 500
    write_workbook(
        out, m,
        [["SPN-PRD-a", long_title, "PASS", "PASS", "JNY-1"]],
        [["UAT-JNY-1", "J", "PASS", "JNY-1", "SPN-PRD-a"]],
        _provenance(),
    )
    ws = openpyxl.load_workbook(out)["REQ"]
    # A 500-character title must not blow the column out arbitrarily wide.
    assert ws.column_dimensions["B"].width <= 50
    # The short-valued verdict columns must still stay readable.
    assert ws.column_dimensions["C"].width >= 10
    assert ws.column_dimensions["D"].width >= 10
